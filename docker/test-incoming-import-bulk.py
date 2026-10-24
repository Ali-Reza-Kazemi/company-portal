"""Automatic order identifiers and atomic spreadsheet/bulk updates."""

import base64
from io import BytesIO
from unittest.mock import patch
from uuid import uuid4

import frappe
import xlsxwriter
from openpyxl import load_workbook
from saeenaria_portal.purchase_reports import incoming_excel as excel
from saeenaria_portal.purchase_reports import incoming_orders as orders
from saeenaria_portal.purchase_reports import product_catalog
from test_support import init_site

init_site()
tokens = []


def content(rows, formula=False):
    stream = BytesIO()
    book = xlsxwriter.Workbook(stream, {"in_memory": True, "strings_to_formulas": False})
    sheet = book.add_worksheet()
    sheet.write_row(0, 0, excel.HEADERS)
    for index, row in enumerate(rows, 1):
        sheet.write_row(index, 0, row)
    if formula:
        sheet.write_formula(1, 2, "=1+1")
    book.close()
    return base64.b64encode(stream.getvalue()).decode()


def preview(rows, **kwargs):
    result = excel.preview(content(rows, **kwargs))
    if result["token"]:
        tokens.append(result["token"])
    return result


def selection(names):
    return [
        {"name": name, "modified": str(frappe.db.get_value(orders.DT, name, "modified"))}
        for name in names
    ]


def rejected(fn, error=frappe.ValidationError, **kwargs):
    try:
        fn(**kwargs)
    except error:
        return
    raise AssertionError("Invalid operation accepted: " + fn.__name__)


try:
    reference = next(
        r.item_code
        for r in frappe.get_all(
            product_catalog.DT,
            filters={"needs_completion": 0, "source_conflict": 0},
            fields=["item_code", "reference_code"],
            limit_page_length=0,
        )
        if r.item_code == r.reference_code
    )
    prefix = "test-auto-" + uuid4().hex
    numeric = int(uuid4().hex[:12], 16)
    first = [prefix, reference, 3, "در راه", "۱۴۰۵/۰۸/۰۲", "آزمون", "توضیح"]
    second = [numeric, reference, 2, "در راه", "1405/08/03", "آزمون", ""]
    with patch.object(frappe, "enqueue") as queue, patch.object(frappe.db, "commit"):
        before = frappe.db.count(orders.DT)
        result = preview([first, second])
        assert not result["errors"] and result["rows"][0]["values"]["expected_date"] == "2026-10-24"
        assert result["rows"][1]["values"]["order_number"] == str(numeric)
        names = [row["values"]["order_key"] for row in result["rows"]]
        assert len(set(names)) == 2 and all(name.startswith("IN-") for name in names)
        assert frappe.db.count(orders.DT) == before
        saved = excel.commit(result["token"])
        assert saved == {"created": 2, "updated": 0} and queue.call_count == 1
        assert excel.commit(result["token"]) == saved and queue.call_count == 1
        # A fresh preview of the same file also matches the existing rows.
        again = preview([first, second])
        assert all(row["action"] == "ویرایش" for row in again["rows"])
        assert [row["values"]["order_key"] for row in again["rows"]] == names
        assert excel.commit(again["token"]) == {"created": 0, "updated": 2}
        assert frappe.db.count(orders.DT) == before + 2
        doc = frappe.get_doc(orders.DT, names[0])
        assert doc.order_number == prefix and str(doc.expected_date) == "2026-10-24"
        orders.save({"name": doc.name, "expected_date_jalali": "۱۴۰۵/۰۸/۰۳"}, str(doc.modified))
        current = selection(names)
        rejected(
            orders.bulk_status,
            frappe.TimestampMismatchError,
            selections=[current[0], {**current[1], "modified": "2000-01-01"}],
            status="Received",
        )
        assert all(frappe.db.get_value(orders.DT, name, "status") == "Active" for name in names)
        assert orders.bulk_status(current, "Received") == {"updated": 2}
        zero = frappe.get_doc(orders.DT, names[1])
        zero.quantity = 0
        zero.save(ignore_permissions=True)
        rejected(orders.bulk_status, selections=selection(names), status="Active")
        assert frappe.db.get_value(orders.DT, names[0], "status") == "Received"
        for rows, options in (
            ([first, first], {}),
            ([first], {"formula": True}),
            ([[*first[:4], "1404/12/30", *first[5:]]], {}),
            ([[*first[:2], -1, *first[3:]]], {}),
            ([[True, *first[1:]]], {}),
        ):
            invalid = preview(rows, **options)
            assert invalid["errors"] and not invalid["token"]
        changed = list(first)
        changed[2] = 4
        successful = preview([changed])
        assert excel.commit(successful["token"]) == {"created": 0, "updated": 1}
        assert frappe.db.get_value(orders.DT, names[0], "quantity") == 4
        update = preview([first])
        doc = frappe.get_doc(orders.DT, names[0])
        doc.notes = "changed"
        doc.save(ignore_permissions=True)
        rejected(excel.commit, frappe.TimestampMismatchError, token=update["token"])
        third = [prefix + "-3", *first[1:]]
        fourth = [prefix + "-4", *first[1:]]
        atomic = preview([third, fourth])
        batch = frappe.cache.get_value(excel.cache_key(atomic["token"]))
        batch["rows"][1]["values"]["quantity"] = "-1"
        frappe.cache.set_value(excel.cache_key(atomic["token"]), batch, expires_in_sec=1800)
        rejected(excel.commit, token=atomic["token"])
        assert not frappe.db.exists(orders.DT, {"order_number": third[0]})
        concurrent = preview([third])
        other = preview([third])
        assert excel.commit(concurrent["token"])["created"] == 1
        rejected(excel.commit, token=other["token"])
        # Legacy identifiers are preserved; ambiguous business keys never guess a row.
        legacy = frappe.get_doc(
            dict(
                doctype=orders.DT,
                order_key="legacy-" + uuid4().hex,
                order_number=prefix,
                reference_code=reference,
                quantity=1,
                status="Active",
            )
        ).insert(ignore_permissions=True)
        assert legacy.order_key.startswith("legacy-")
        assert preview([first])["errors"]
        excel.template()
        book = load_workbook(BytesIO(frappe.local.response.filecontent))
        assert [cell.value for cell in book.active[1]] == excel.HEADERS
        assert "شناسه ردیف سفارش" not in excel.HEADERS and len(excel.HEADERS) == 7
        book.close()
        frappe.set_user("Guest")
        for fn, args in (
            (excel.preview, {"content": ""}),
            (excel.commit, {"token": result["token"]}),
            (excel.template, {}),
            (orders.bulk_status, {"selections": current, "status": "Received"}),
        ):
            rejected(fn, frappe.PermissionError, **args)
    print(
        "PASS automatic identifiers, legacy preservation, Excel without identifiers, fresh-file deduplication, ambiguous/concurrent import rejection, dates, numeric codes, stale locks, atomic bulk/import and permissions"
    )
finally:
    frappe.db.rollback()
    for token in tokens:
        frappe.cache.delete_value(excel.cache_key(token))
    frappe.destroy()
