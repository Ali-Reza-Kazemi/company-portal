"""Transactional integration checks for workbook I/O and recurrence; no production counts retained."""
import os,base64,json
from io import BytesIO
from unittest.mock import patch
from types import SimpleNamespace
import jdatetime
import frappe
from test_support import init_site
init_site()
from openpyxl import load_workbook
from saeenaria_portal.purchase_reports import inventory_excel as x
from saeenaria_portal.schedule import nominal_dates
from saeenaria_portal.testing import portal_fixtures
preview=None
try:
 with portal_fixtures() as fixture, patch.dict(frappe.conf, {'inventory_reporting_owner':None}):
  if not frappe.db.exists('Portal Domain','inventory'):
   frappe.get_doc(dict(doctype='Portal Domain',domain_code='inventory',title='Test inventory')).insert()
  frappe.conf.inventory_reporting_owner=fixture.employee
  t=SimpleNamespace(frequency='Semiannual',day_of_month=15)
  dates=list(nominal_dates(t,jdatetime.date(1406,1,1).togregorian(),jdatetime.date(1406,12,29).togregorian()))
  assert [jdatetime.date.fromgregorian(date=d).strftime('%m/%d') for d in dates]==['01/15','07/15']
  safe=x.workbook(['text'],[['=1+1']],'test');b=load_workbook(BytesIO(safe));assert b.active['A2'].data_type=='s'
  raw=x.workbook(x.HEADERS,[['000123',1,18],['000124',1,0]],'counts');content=base64.b64encode(raw).decode()
  cash=SimpleNamespace(name='TEST-CASH',domain='inventory',kind='Cash',payload=json.dumps({'end':'2099-01-01','urgency':[{'code':'000123','title':'Test item','priority':'Test','quantity':5,'value_rial':100}]}))
  with patch.object(x.frappe,'get_doc',return_value=cash):x.export_urgency(cash.name)
  exported=load_workbook(BytesIO(frappe.local.response.filecontent),read_only=True)
  expected=json.loads(cash.payload)['urgency']
  assert exported.active.max_row==len(expected)+1
  assert exported.active['A2'].value==str(expected[0]['code'])
  assert exported.active['D2'].data_type=='n'
  exported.close()
  count_before=frappe.db.count('Portal Record')
  frappe.db.savepoint('excel_counts')
  with patch.object(x.w.source,'current_book',return_value=[{'quantity':20}]):
   preview=x.preview_counts(content);assert preview['token'] and not preview['errors'] and len(preview['rows'])==2
   with patch.object(frappe.db,'commit'):
    result=x.commit_counts(preview['token']);assert result['created']==2
    again=x.commit_counts(preview['token']);assert again['created']==0 and again['skipped']==2
   rows=frappe.get_all('Portal Record',filters={'external_id':['like','inventory-count-import:%']},fields=['quantity','baseline_snapshot'])
   assert any(r.quantity==18 and json.loads(r.baseline_snapshot)['book']=='20' for r in rows)
  frappe.db.rollback(save_point='excel_counts');assert frappe.db.count('Portal Record')==count_before
  with patch.object(x.w.source,'current_book',return_value=[{'quantity':21}]):
   try:x.commit_counts(preview['token']);raise AssertionError('changed baseline accepted')
   except frappe.ValidationError:pass
  duplicate=x.workbook(x.HEADERS,[['000123',1,18],['000123',1,17]],'counts')
  with patch.object(x.w.source,'current_book',return_value=[{'quantity':20}]):
   result=x.preview_counts(base64.b64encode(duplicate).decode());assert not result['token'] and result['errors'][0]['row']==3
  frappe.set_user('Guest')
  try:x.count_template();raise AssertionError('guest accepted')
  except frappe.PermissionError:pass
  print('PASS: semiannual dates, real XLSX, formula-safe text, leading zero codes, count preview, atomic creation, retry deduplication, baseline change rejection, duplicate row rejection, guest permissions.')
finally:
 if preview and preview.get('token'):frappe.cache.delete_value('count-import:'+preview['token'])
 frappe.db.rollback();frappe.destroy()
