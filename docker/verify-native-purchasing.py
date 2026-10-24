"""Read-only authenticated HTTP verification of the active native purchase UI/API."""

import base64
import http.cookiejar
import json
import os
import re
import urllib.parse
import urllib.request

import frappe
from saeenaria_portal.purchase_reports import service
from test_support import init_site

init_site()

client = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))
base = os.environ.get("PORTAL_TEST_HTTP_URL", "http://127.0.0.1:8000")
csrf = ""
preview_token = None


def request(path, data=None):
    headers = {"Host": frappe.local.site}
    if data is not None:
        headers["X-Frappe-CSRF-Token"] = csrf
        data = urllib.parse.urlencode(data).encode()
    return client.open(
        urllib.request.Request(base + path, data=data, headers=headers), timeout=60
    ).read()


def api(method, **values):
    return json.loads(
        request(
            "/api/method/saeenaria_portal.purchase_reports."
            + method
            + "?"
            + urllib.parse.urlencode(values)
        )
    )["message"]


try:
    request("/api/method/login", {"usr": "Administrator", "pwd": os.environ["ADMIN_PASSWORD"]})
    html = request("/portal-purchase").decode()
    csrf = re.search(r'<meta name="csrf-token" content="([^"]+)"', html).group(1)
    assert 'id="purchase-incoming-panel"' in html and 'id="incoming-form"' in html
    assert 'id="incoming-file"' in html and 'id="incoming-bulk-apply"' in html
    assert "data-jalali-date" in html and "کد سفارش" in html
    assert 'id="incoming-order_key"' not in html
    assert "تاریخ مورد انتظار (میلادی)" not in html
    assert "portal-dates.js?v=4" in html
    from io import BytesIO

    from openpyxl import load_workbook

    template = load_workbook(
        BytesIO(request("/api/method/saeenaria_portal.purchase_reports.incoming_excel.template"))
    )
    assert template.active.cell(1, 1).value == "کد سفارش"
    template.close()
    import xlsxwriter
    from saeenaria_portal.purchase_reports import incoming_excel, product_catalog

    assert frappe.get_meta("Portal Incoming Order").get_field("order_number").label == "کد سفارش"
    reference = next(
        row.item_code
        for row in frappe.get_all(
            product_catalog.DT,
            filters={"needs_completion": 0, "source_conflict": 0},
            fields=["item_code", "reference_code"],
            limit_page_length=0,
        )
        if row.item_code == row.reference_code
    )
    stream = BytesIO()
    workbook = xlsxwriter.Workbook(stream, {"in_memory": True})
    sheet = workbook.add_worksheet()
    sheet.write_row(0, 0, incoming_excel.HEADERS)
    sheet.write_row(1, 0, ["PO-http-preview-only", reference, 2, "در راه", "۱۴۰۵/۰۸/۰۲", "", ""])
    workbook.close()
    preview = json.loads(
        request(
            "/api/method/saeenaria_portal.purchase_reports.incoming_excel.preview",
            {"content": base64.b64encode(stream.getvalue()).decode()},
        )
    )["message"]
    preview_token = preview["token"]
    assert not preview["errors"] and preview_token
    assert preview["rows"][0]["values"]["expected_date"] == "2026-10-24"

    assert "purchase-incoming.js?v=3" in html and "purchase-reports.js?v=8" in html
    js = request("/assets/saeenaria_portal/js/purchase-reports.js?v=8").decode()
    assert "location.replace('/portal-products?status=issues')" in js
    assert "PurchaseIncoming.load" in js
    assert (
        "incoming_orders."
        in request("/assets/saeenaria_portal/js/purchase-incoming.js?v=3").decode()
    )
    assert 'value="issues"' in request("/portal-products?status=issues").decode()
    overview = api("service.overview")
    assert len(overview["reports"]) == 7 and all(
        r["name"] != "reconcile" for r in overview["reports"]
    )
    assert not overview["state"]["errors"]
    assert api("incoming_orders.overview")["can_edit"]
    assert api("product_catalog.listing", status="issues")["matched"] >= 0
    counts = {}
    for name in service.REPORTS:
        result = api("service.report", name=name)
        assert result["available"] and result["settings_at"][0]["source"] == "portal-native-v2"
        counts[name] = result["total"]
    assert service.read_snapshot("combined")["processor"] == "portal-native-v2"
    print(json.dumps({"native_http": "passed", "report_totals": counts}, ensure_ascii=False))
finally:
    if preview_token:
        frappe.cache.delete_value("incoming-import:" + preview_token)
    if csrf:
        request("/api/method/logout", {})
    frappe.destroy()
