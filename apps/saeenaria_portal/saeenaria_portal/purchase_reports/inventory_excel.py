"""Native Excel downloads and previewed, atomic physical-count imports."""
import base64
import hashlib
import json
from io import BytesIO
import math
import uuid
import zipfile

import frappe
from frappe.utils import now_datetime, today
from . import workflow as w

HEADERS = ['کد کالا', 'شناسه انبار', 'تعداد شمارش‌شده']


def workbook(headers, rows, title, text_columns=None):
    import xlsxwriter
    stream = BytesIO()
    book = xlsxwriter.Workbook(stream, {'in_memory': True, 'strings_to_formulas': False, 'strings_to_urls': False})
    sheet = book.add_worksheet(title)
    sheet.right_to_left()
    heading = book.add_format({'bold': True, 'bg_color': '#173DA6', 'font_color': '#FFFFFF', 'text_wrap': True})
    sheet.write_row(0, 0, headers, heading)
    for index, row in enumerate(rows, 1):
        sheet.write_row(index, 0, row)
    sheet.freeze_panes(1, 0)
    sheet.autofilter(0, 0, len(rows), len(headers)-1)
    sheet.set_column(0, len(headers)-1, 23)
    for column in ([0] if text_columns is None else text_columns):
        sheet.set_column(column, column, 23, book.add_format({'num_format': '@'}))
    sheet.set_row(0, 32)
    book.close()
    return stream.getvalue()


def download(content, filename):
    frappe.local.response.filename = filename
    frappe.local.response.filecontent = content
    frappe.local.response.type = 'binary'
    frappe.local.response.display_content_as = 'attachment'


@frappe.whitelist()
def export_urgency(snapshot):
    w.require_access()
    doc = frappe.get_doc('Portal Inventory Snapshot', snapshot)
    if doc.domain != 'inventory' or doc.kind != 'Cash':
        frappe.throw('نسخه گزارش نقدینگی معتبر لازم است')
    data = json.loads(doc.payload)
    numeric = lambda value: None if value is None else float(value)
    rows = [[str(r['code']), r.get('title'), r.get('priority'), numeric(r.get('quantity')), r.get('order_date'), numeric(r.get('coverage_days')), numeric(r.get('lead_days')), None if r.get('value_rial') is None else float(r['value_rial'])/10, str(data.get('planner_asof') or data['end']), doc.name] for r in data.get('urgency', [])]
    download(workbook(['کد کالا','نام کالا','فوریت','تعداد سفارش','تاریخ پیشنهادی میلادی','پوشش (روز)','زمان تأمین (روز)','برآورد (تومان)','تاریخ مبنا میلادی','نسخه گزارش'], rows, 'زمان‌بندی خرید'), 'purchase-planning.xlsx')


@frappe.whitelist()
def count_template():
    w.require_input()
    # No example quantities: an untouched template cannot create fictitious counts.
    download(workbook(HEADERS, [], 'شمارش انبار'), 'stock-count-template.xlsx')


@frappe.whitelist(methods=['POST'])
def preview_counts(content):
    w.require_input()
    from openpyxl import load_workbook
    if not isinstance(content, str) or len(content) > 7_000_000:
        frappe.throw('حداکثر اندازه فایل اکسل ۵ مگابایت است')
    try:
        raw = base64.b64decode(content, validate=True)
        with zipfile.ZipFile(BytesIO(raw)) as archive:
            if sum(f.file_size for f in archive.infolist()) > 30_000_000:
                frappe.throw('حجم محتوای فایل بیش از حد مجاز است')
        book = load_workbook(BytesIO(raw), read_only=True, data_only=False)
    except Exception:
        frappe.throw('فایل XLSX معتبر نیست؛ از قالب همین صفحه استفاده کنید')
    sheet = book.active
    if sheet.max_row > 1001 or sheet.max_column > 3:
        book.close()
        frappe.throw('حداکثر ۱۰۰۰ ردیف و سه ستون قالب مجاز است')
    source = sheet.iter_rows()
    headers = next(source, [])
    if [c.value for c in headers] != HEADERS:
        book.close()
        frappe.throw('عنوان ستون‌ها باید مطابق قالب دانلودشده باشد')
    rows, errors, seen = [], [], set()
    for index, cells in enumerate(source, 2):
        if all(c.value is None for c in cells):
            continue
        try:
            if any(c.data_type == 'f' for c in cells):
                raise ValueError('فرمول مجاز نیست')
            code, stock, quantity = [c.value for c in cells]
            if not isinstance(code, str) or not code.strip():
                raise ValueError('کد کالا باید متن باشد تا صفرهای ابتدای آن حفظ شوند')
            code = code.strip()
            if isinstance(stock, bool) or stock is None or float(stock) != int(stock):
                raise ValueError('شناسه انبار باید عدد صحیح باشد')
            stock = int(stock)
            if isinstance(quantity, bool) or quantity is None or not math.isfinite(float(quantity)) or float(quantity) < 0:
                raise ValueError('تعداد شمارش باید عدد نامنفی باشد')
            if (code, stock) in seen:
                raise ValueError('کالا و انبار در فایل تکرار شده است')
            seen.add((code, stock))
            current = w.source.current_book(code, stock)
            if not current or current[0]['quantity'] is None:
                raise ValueError('کالا یا موجودی این انبار در سپیدار یافت نشد')
            rows.append({'row': index, 'code': code, 'stock': stock, 'quantity': float(quantity), 'book': str(current[0]['quantity'])})
        except Exception as exc:
            errors.append({'row': index, 'message': str(exc)[:200]})
    book.close()
    if not rows and not errors:
        frappe.throw('فایل شمارشی ندارد')
    token = uuid.uuid4().hex
    if not errors:
        frappe.cache.set_value('count-import:'+token, {'user': frappe.session.user, 'rows': rows, 'hash': hashlib.sha256(raw).hexdigest()}, expires_in_sec=1800)
    return {'token': token if not errors else None, 'rows': rows, 'errors': errors}


@frappe.whitelist(methods=['POST'])
def commit_counts(token):
    w.require_input()
    batch = frappe.cache.get_value('count-import:'+str(token))
    if not batch or batch['user'] != frappe.session.user:
        frappe.throw('پیش‌نمایش منقضی شده؛ فایل را دوباره بررسی کنید')
    created, skipped = 0, 0
    with frappe.cache.lock('count-import-lock:'+batch['hash'], timeout=120):
        frappe.db.savepoint('count_import')
        try:
            for row in batch['rows']:
                external = 'inventory-count-import:'+batch['hash']+':'+str(row['row'])
                if frappe.db.exists('Portal Record', {'external_id': external}):
                    skipped += 1
                    continue
                current = w.source.current_book(row['code'], row['stock'])
                if not current or str(current[0]['quantity']) != row['book']:
                    frappe.throw('موجودی دفتری از زمان پیش‌نمایش تغییر کرده؛ فایل را دوباره بررسی کنید')
                info = {'policy': w.math.POLICY, 'stock': row['stock'], 'book': row['book'], 'book_captured_at': str(now_datetime()), 'import_hash': batch['hash'], 'import_row': row['row']}
                frappe.get_doc(dict(doctype='Portal Record', domain='inventory', title='شمارش موجودی: '+row['code'], record_type='Stock Count', external_id=external, record_date=today(), responsible=w.owner(), channel='inventory-monitoring', item_code=row['code'], quantity=row['quantity'], locked=1, status='Counted', baseline_snapshot=json.dumps(info), is_demo=0)).insert(ignore_permissions=True)
                created += 1
            # Commit while holding the batch lock, so concurrent retries cannot duplicate rows.
            frappe.db.commit()
        except Exception:
            frappe.db.rollback(save_point='count_import')
            raise
    return {'created': created, 'skipped': skipped}
