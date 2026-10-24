"""Optional XLSX ingestion into native orders, with a versioned preview and atomic save."""

import base64
import re
import zipfile
from decimal import Decimal, InvalidOperation
from io import BytesIO
from uuid import uuid4

import frappe
import jdatetime
from frappe.utils import get_datetime

from . import incoming_orders as orders
from .product_catalog import require_edit

HEADERS = [
    "کد سفارش",
    "کد کالای مرجع",
    "تعداد باقی‌مانده در راه",
    "وضعیت",
    "تاریخ مورد انتظار (شمسی)",
    "تأمین‌کننده",
    "توضیحات",
]
FIELDS = [
    "order_number",
    "reference_code",
    "quantity",
    "status",
    "expected_date",
    "supplier",
    "notes",
]
STATUSES = {
    "در راه": "Active",
    "دریافت‌شده": "Received",
    "لغوشده": "Cancelled",
    "Active": "Active",
    "Received": "Received",
    "Cancelled": "Cancelled",
}


def jalali_date(value):
    if value in (None, ""):
        return None
    if not isinstance(value, str):
        frappe.throw("تاریخ باید متن شمسی به صورت ۱۴۰۵/۰۸/۰۲ باشد")
    value = value.strip().translate(str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789"))
    if not re.fullmatch(r"[0-9]{4}/[0-9]{1,2}/[0-9]{1,2}", value):
        frappe.throw("تاریخ شمسی باید به صورت سال/ماه/روز باشد")
    try:
        year, month, day = map(int, value.split("/"))
        if not 1300 <= year <= 1500:
            raise ValueError()
        return jdatetime.date(year, month, day).togregorian().isoformat()
    except ValueError:
        frappe.throw("تاریخ شمسی معتبر نیست")


@frappe.whitelist()
def template():
    require_edit()
    from .inventory_excel import download, workbook

    download(
        workbook(HEADERS, [], "سفارش‌های در راه", text_columns=[0, 1, 3, 4, 5, 6]),
        "incoming-orders-template.xlsx",
    )


def matching_orders(values):
    return frappe.get_all(
        orders.DT,
        filters={
            "order_number": values["order_number"],
            "reference_code": values["reference_code"],
        },
        fields=["name"],
        limit_page_length=3,
    )


def cache_key(token):
    return "incoming-import:" + str(token)


@frappe.whitelist(methods=["POST"])
def preview(content):
    require_edit()
    from openpyxl import load_workbook

    if not isinstance(content, str) or len(content) > 7_000_000:
        frappe.throw("حداکثر اندازه فایل XLSX پنج مگابایت است")
    try:
        raw = base64.b64decode(content, validate=True)
        if len(raw) > 5_000_000:
            raise ValueError()
        with zipfile.ZipFile(BytesIO(raw)) as archive:
            if (
                sum(info.file_size for info in archive.infolist()) > 30_000_000
                or len(archive.infolist()) > 1000
            ):
                raise ValueError()
        book = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    except Exception:
        frappe.throw("فایل XLSX معتبر نیست یا حجم آن بیش از حد مجاز است")
    result, errors, seen = [], [], set()
    try:
        sheet = book.active
        if sheet.max_row > 1001 or sheet.max_column != len(HEADERS):
            frappe.throw("حداکثر ۱۰۰۰ ردیف و ستون‌های قالب مجاز است")
        source = sheet.iter_rows()
        if [cell.value for cell in next(source, [])] != HEADERS:
            frappe.throw("عنوان ستون‌ها باید مطابق قالب دانلودشده باشد")
        for index, cells in enumerate(source, 2):
            if index > 1001:
                frappe.throw("حداکثر ۱۰۰۰ ردیف مجاز است")
            if all(cell.value is None for cell in cells):
                continue
            try:
                if any(cell.data_type in ("f", "e") for cell in cells):
                    frappe.throw("فرمول و خطای اکسل مجاز نیست")
                values = dict(zip(FIELDS, [cell.value for cell in cells], strict=True))
                if isinstance(values["order_number"], (int, float)) and not isinstance(
                    values["order_number"], bool
                ):
                    numeric_code = Decimal(str(values["order_number"]))
                    if numeric_code.is_finite():
                        values["order_number"] = (
                            str(int(numeric_code))
                            if numeric_code == numeric_code.to_integral_value()
                            else str(numeric_code)
                        )
                for field in ("order_number", "reference_code"):
                    if not isinstance(values[field], str) or not values[field].strip():
                        frappe.throw(
                            "کد سفارش و کد مرجع باید متن غیرخالی باشند؛ قالب سلول را Text بگذارید"
                        )
                    values[field] = values[field].strip()
                values["status"] = STATUSES.get(str(values["status"] or "در راه").strip())
                if not values["status"]:
                    frappe.throw("وضعیت باید در راه، دریافت‌شده یا لغوشده باشد")
                try:
                    qty = Decimal(str(values["quantity"]))
                    if (
                        isinstance(values["quantity"], bool)
                        or not qty.is_finite()
                        or qty < 0
                        or qty > 1_000_000_000
                        or qty.as_tuple().exponent < -9
                    ):
                        raise ValueError()
                except (ValueError, TypeError, InvalidOperation):
                    frappe.throw("تعداد باقی‌مانده نامعتبر است")
                values["quantity"] = str(qty)
                values["expected_date"] = jalali_date(values["expected_date"])
                identity = (values["order_number"].casefold(), values["reference_code"].casefold())
                if identity in seen:
                    frappe.throw(
                        "کد سفارش و کد مرجع در فایل تکراری است؛ هر کالا را در هر سفارش یک بار وارد کنید"
                    )
                seen.add(identity)
                matches = matching_orders(values)
                if len(matches) > 1:
                    frappe.throw(
                        "برای این کد سفارش و کالا چند ردیف موجود است؛ آن‌ها را از صفحه سفارش‌ها ویرایش کنید"
                    )
                name = matches[0].name if matches else None
                doc = frappe.get_doc(orders.DT, name) if name else frappe.new_doc(orders.DT)
                modified = str(doc.modified) if name else None
                doc.update(values)
                orders.validate_document(doc)
                values["order_key"] = doc.order_key
                values["reference_code"] = doc.reference_code
                result.append(
                    {
                        "row": index,
                        "name": name,
                        "modified": modified,
                        "values": values,
                        "action": "ویرایش" if name else "جدید",
                    }
                )
            except frappe.ValidationError as exc:
                errors.append({"row": index, "message": str(exc)[:300]})
    finally:
        book.close()
    if not result and not errors:
        frappe.throw("فایل ردیفی برای ورود ندارد")
    token = uuid4().hex if not errors else None
    if token:
        frappe.cache.set_value(
            cache_key(token),
            {"version": 2, "user": frappe.session.user, "rows": result},
            expires_in_sec=1800,
        )
    return {"token": token, "rows": result, "errors": errors}


@frappe.whitelist(methods=["POST"])
def commit(token):
    require_edit()
    with frappe.cache.lock("incoming-import-write:" + frappe.local.site, timeout=120):
        batch = frappe.cache.get_value(cache_key(token))
        if not batch or batch.get("version") != 2 or batch["user"] != frappe.session.user:
            frappe.throw("پیش‌نمایش منقضی شده؛ فایل را دوباره بررسی کنید")
        if batch.get("result"):
            return batch["result"]
        frappe.db.savepoint("incoming_import")
        created = updated = 0
        try:
            for row in sorted(batch["rows"], key=lambda row: row["values"]["order_key"]):
                matches = matching_orders(row["values"])
                if {item.name for item in matches} != ({row["name"]} if row["name"] else set()):
                    frappe.throw(
                        "ردیف‌های این کد سفارش و کالا پس از پیش‌نمایش تغییر کرده‌اند؛ فایل را دوباره بررسی کنید"
                    )
                if row["name"]:
                    frappe.db.get_value(orders.DT, row["name"], "name", for_update=True)
                    doc = frappe.get_doc(orders.DT, row["name"])
                    if get_datetime(doc.modified) != get_datetime(row["modified"]):
                        frappe.throw(
                            "سفارش از زمان پیش‌نمایش تغییر کرده است؛ فایل را دوباره بررسی کنید",
                            frappe.TimestampMismatchError,
                        )
                else:
                    if frappe.db.exists(orders.DT, {"order_key": row["values"]["order_key"]}):
                        frappe.throw(
                            "ردیف سفارش پس از پیش‌نمایش ثبت شده است؛ فایل را دوباره بررسی کنید"
                        )
                    doc = frappe.new_doc(orders.DT)
                doc.update(row["values"])
                doc.save(ignore_permissions=True)
                if row["name"]:
                    updated += 1
                else:
                    created += 1
            orders.enqueue_recalculation()
            frappe.db.commit()
        except Exception:
            frappe.db.rollback(save_point="incoming_import")
            raise
        batch["result"] = {"created": created, "updated": updated}
        frappe.cache.set_value(cache_key(token), batch, expires_in_sec=1800)
        return batch["result"]
